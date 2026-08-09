"""Getting text out of an uploaded file (section 4.3, pipeline step 1).

Deliberately separate from the LLM. This module answers "what words are in this
file", which is a mechanical question with a right answer; :mod:`app.services.llm`
answers "what do those words require of this client", which is not. Keeping them
apart means a bad extraction is diagnosable — the parsed text is stored, so
"the model got it wrong" and "the model never saw the text" are distinguishable
at a glance.

Each extractor returns ``(text, method)``. The method is recorded on the
document because it calibrates trust: text lifted from a PDF's text layer is
exact, and text from a scan is a guess. The UI says which one it has.

**Scanned documents.** A PDF whose pages carry no text layer is a scan. There
is no OCR engine in this deployment, so rather than returning an empty string
that would read as "an empty notice", extraction reports the condition and the
caller routes the file to the vision model instead.
"""
from __future__ import annotations

import logging
import re
from pathlib import Path

logger = logging.getLogger(__name__)

# Below this many characters, a PDF's "text layer" is page furniture — a header,
# a page number — rather than content. Treated as a scan.
MIN_TEXT_LAYER_CHARS = 40

# How much text is handed to a model. A GST circular runs to a few thousand
# words; a consolidated Finance Act does not fit any free-tier context window.
# Truncating at a bounded size with a marker is honest, where silently sending
# the first N bytes is not.
MAX_MODEL_CHARS = 24_000

_WHITESPACE = re.compile(r"[ \t\xa0]+")
_BLANK_LINES = re.compile(r"\n{3,}")


class ExtractionError(Exception):
    """The file could not be read at all."""


class ScannedDocument(ExtractionError):
    """The file is an image-only PDF. The caller should try the vision model."""


def normalise(text: str) -> str:
    """Collapse the whitespace damage that PDF extraction always produces.

    Not cosmetic. A model charged per token is handed roughly a third fewer of
    them once the runs of spaces a two-column layout produces are collapsed,
    and the extracted text is also what a human reads when checking the
    model's answer.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _WHITESPACE.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return _BLANK_LINES.sub("\n\n", text).strip()


def _extract_pdf(path: Path) -> tuple[str, str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001 - one bad page must not lose the rest
            logger.warning("Page extraction failed in %s: %s", path.name, exc)
            pages.append("")

    text = normalise("\n\n".join(pages))
    if len(text) < MIN_TEXT_LAYER_CHARS:
        raise ScannedDocument(
            f"{path.name} has no usable text layer across {len(reader.pages)} page(s)"
        )
    return text, "pdf_text"


def _extract_docx(path: Path) -> tuple[str, str]:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]

    # Tables carry the substance in a filing annexure, and python-docx does not
    # include them in ``paragraphs``. Omitting them silently drops exactly the
    # figures the extraction exists to find.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return normalise("\n".join(parts)), "docx"


def _extract_html(path: Path) -> tuple[str, str]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "lxml")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    return normalise(soup.get_text("\n")), "html"


def _extract_text(path: Path) -> tuple[str, str]:
    return normalise(path.read_text(encoding="utf-8", errors="replace")), "text"


_BY_SUFFIX = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".doc": _extract_docx,
    ".html": _extract_html,
    ".htm": _extract_html,
    ".txt": _extract_text,
    ".md": _extract_text,
    ".csv": _extract_text,
}


def extract_text(path: str | Path, *, mime_type: str | None = None) -> tuple[str, str]:
    """``(text, method)`` for a file on disk.

    Dispatches on the file's extension, falling back to the MIME type the
    upload declared. The extension is preferred because it describes the bytes
    that were stored, while the MIME type is whatever the client claimed.

    Raises :class:`ScannedDocument` for an image-only PDF and
    :class:`ExtractionError` for anything unreadable.
    """
    file_path = Path(path)
    if not file_path.exists():
        raise ExtractionError(f"No file at {file_path}")

    extractor = _BY_SUFFIX.get(file_path.suffix.lower())

    if extractor is None and mime_type:
        _DOCX_MIME = (
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document"
        )
        by_mime = {
            "application/pdf": _extract_pdf,
            _DOCX_MIME: _extract_docx,
            "text/html": _extract_html,
            "text/plain": _extract_text,
            "text/csv": _extract_text,
        }
        extractor = by_mime.get(mime_type.split(";")[0].strip().lower())

    if extractor is None:
        raise ExtractionError(
            f"No extractor for {file_path.suffix or mime_type or 'unknown type'}"
        )

    try:
        return extractor(file_path)
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 - a corrupt upload is not a crash
        raise ExtractionError(f"Could not read {file_path.name}: {exc}") from exc


def clip_for_model(text: str, *, limit: int = MAX_MODEL_CHARS) -> str:
    """Trim text to what a model will accept, saying so where it was cut.

    The marker matters: without it a model reading a truncated circular has no
    way to know it is missing the operative paragraph, and will answer as
    confidently about the half it received as about the whole.
    """
    if len(text) <= limit:
        return text
    return text[:limit] + f"\n\n[TRUNCATED — {len(text) - limit} more characters omitted]"
